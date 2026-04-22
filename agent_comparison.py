"""
Compare 3 evidence-gathering strategies on ChartCheck test set:

  1. Visual Only     — Agent 2 alone  (reads chart image directly)
  2. Structured Only — Agent 3 alone  (extracts table first, then reasons)
  3. Combined        — Agent 2 + Agent 3 + resolution logic (Agent 4):
                         • Both agree AND avg confidence >= 0.70
                           → accept shared verdict directly
                         • Disagree OR low confidence
                           → classify claim type:
                               numerical → trust Structured (Agent 3)
                               visual    → trust Visual (Agent 2)

Agents 2 & 3 always run in parallel (same as production system).

Usage:
  python agent_comparison.py
  python agent_comparison.py --model openrouter/google/gemini-2.5-flash-preview
  python agent_comparison.py --limit 50 --split test
  python agent_comparison.py --threshold 0.80
"""

import argparse
import asyncio
import json
import time
from datetime import datetime
from pathlib import Path

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from dotenv import load_dotenv

load_dotenv(Path(__file__).parent / "chart_verifier" / ".env")

from benchmark_eval import (
    load_samples, download_image, encode_image, remove_temp,
    call_llm, run_visual, run_structured,
    candidate_to_verdict, compute_metrics, print_metrics,
)


def extract_verdict(result: dict | None) -> tuple[str, float]:
    """
    Extract (TRUE/FALSE/UNKNOWN, confidence) from an agent result dict.

    Handles both output formats:
      New agents: {"verdict": "correct | incorrect", "confidence": ...}
      Old agents: {"candidate_answer": "supported | contradicted | ...", "confidence": ...}
    """
    if not result:
        return "UNKNOWN", 0.0

    conf = float(result.get("confidence") or 0.0)

    # New format — verdict: correct | incorrect
    verdict_field = (result.get("verdict") or "").lower().strip()
    if verdict_field == "correct":
        return "TRUE", conf
    if verdict_field == "incorrect":
        return "FALSE", conf

    # Old format — candidate_answer: supported | contradicted | ...
    return candidate_to_verdict(result.get("candidate_answer")), conf

# ── Configure ─────────────────────────────────────────────────────────────────

MODEL                = "openrouter/anthropic/claude-sonnet-4.6"
LIMIT                = 30
SPLIT                = "test"
CONFIDENCE_THRESHOLD = 0.70

OUTPUT_DIR = Path(__file__).parent / "output"

# ── Claim type classifier ─────────────────────────────────────────────────────

CLASSIFY_INSTRUCTION = """
Classify the following claim as either "numerical" or "visual".

numerical — the claim mentions specific numbers, percentages, quantities,
            or makes a precise quantity comparison.
            Examples: "only 12.08% of...", "more than 50 stores", "increased by 3%"

visual    — the claim is about a trend, direction, ranking, pattern, or appearance
            without citing specific measured values.
            Examples: "the line increases", "the largest slice", "in the lowest place"

Reply with ONLY one word: numerical  or  visual
"""


async def classify_claim(claim: str, model: str) -> str:
    raw = await call_llm(
        model,
        CLASSIFY_INSTRUCTION,
        [{"type": "text", "text": f"Claim: {claim}"}],
    )
    return "numerical" if "numerical" in raw.lower() else "visual"


# ── Resolution logic (Agent 4) ────────────────────────────────────────────────

async def resolve(
    claim: str,
    vis: dict | None,
    strc: dict | None,
    model: str,
    threshold: float,
) -> tuple[str, str]:
    """
    Returns (final_verdict, resolution_reason).
    Mirrors the logic in verdict_feedback.py.
    """
    vis_pred,  vis_conf  = extract_verdict(vis)
    strc_pred, strc_conf = extract_verdict(strc)
    avg_conf = (vis_conf + strc_conf) / 2.0

    agents_agree = vis_pred == strc_pred and vis_pred != "UNKNOWN"

    # Case A — both agree with high confidence: no tiebreaker needed
    if agents_agree and avg_conf >= threshold:
        return vis_pred, f"agree+conf={avg_conf:.2f} → accept directly"

    # Case B — disagree or low confidence: classify then pick winner
    claim_type = await classify_claim(claim, model)
    if claim_type == "numerical":
        verdict = strc_pred if strc_pred != "UNKNOWN" else vis_pred
        return verdict, f"{'disagree' if not agents_agree else 'low-conf'} → numerical → trust structured (conf={strc_conf:.2f})"
    else:
        verdict = vis_pred if vis_pred != "UNKNOWN" else strc_pred
        return verdict, f"{'disagree' if not agents_agree else 'low-conf'} → visual → trust visual (conf={vis_conf:.2f})"


# ── Disagreement analysis ─────────────────────────────────────────────────────

def disagreement_analysis(rows: list[dict]) -> dict:
    counts = {
        "all_correct":          0,
        "combined_wins_both":   0,
        "combined_wins_visual": 0,
        "combined_wins_struct": 0,
        "visual_only_wins":     0,
        "struct_only_wins":     0,
        "all_wrong":            0,
    }
    for r in rows:
        gt   = r["ground_truth"]
        v_ok = r["visual_verdict"]     == gt
        s_ok = r["structured_verdict"] == gt
        c_ok = r["combined_verdict"]   == gt

        if v_ok and s_ok and c_ok:
            counts["all_correct"] += 1
        elif c_ok and not v_ok and not s_ok:
            counts["combined_wins_both"] += 1
        elif c_ok and not v_ok and s_ok:
            counts["combined_wins_visual"] += 1
        elif c_ok and v_ok and not s_ok:
            counts["combined_wins_struct"] += 1
        elif v_ok and not s_ok and not c_ok:
            counts["visual_only_wins"] += 1
        elif s_ok and not v_ok and not c_ok:
            counts["struct_only_wins"] += 1
        elif not v_ok and not s_ok and not c_ok:
            counts["all_wrong"] += 1

    return counts


# ── Excel output ──────────────────────────────────────────────────────────────

def save_excel(rows: list[dict], metrics: dict, disagree: dict, path: Path) -> None:
    wb     = openpyxl.Workbook()
    hf     = Font(bold=True, color="FFFFFF")
    hfill  = PatternFill("solid", fgColor="2E4057")
    green  = PatternFill("solid", fgColor="C6EFCE")
    red    = PatternFill("solid", fgColor="FFC7CE")
    yellow = PatternFill("solid", fgColor="FFEB9C")
    center = Alignment(horizontal="center")

    def style_header(ws, row_idx=1):
        for cell in ws[row_idx]:
            cell.font = hf; cell.fill = hfill; cell.alignment = center

    def autofit(ws):
        for col in ws.columns:
            w = max((len(str(c.value or "")) for c in col), default=10)
            ws.column_dimensions[get_column_letter(col[0].column)].width = min(w + 4, 60)

    # ── Sheet 1: Summary ──────────────────────────────────────────────────────
    ws1 = wb.active
    ws1.title = "Summary"
    ws1.append(["Approach", "Accuracy", "Precision", "Recall", "F1",
                "TP", "FP", "TN", "FN", "Unknown", "Total"])
    style_header(ws1)
    labels = [("Visual Only",     "visual_only"),
              ("Structured Only", "structured_only"),
              ("Combined",        "combined")]
    for label, key in labels:
        m = metrics[key]
        ws1.append([label,
                    round(m["accuracy"],  3), round(m["precision"], 3),
                    round(m["recall"],    3), round(m["f1"],        3),
                    m["tp"], m["fp"], m["tn"], m["fn"], m["unknown"], m["total"]])
    autofit(ws1)

    # ── Sheet 2: Per-sample ───────────────────────────────────────────────────
    ws2 = wb.create_sheet("Per_Sample")
    ws2.append(["#", "Claim", "Ground Truth",
                "Visual", "V-OK", "V-Conf",
                "Structured", "S-OK", "S-Conf",
                "Combined", "C-OK", "Resolution"])
    style_header(ws2)
    for i, r in enumerate(rows, 1):
        gt   = r["ground_truth"]
        v_ok = r["visual_verdict"]     == gt
        s_ok = r["structured_verdict"] == gt
        c_ok = r["combined_verdict"]   == gt
        ws2.append([
            i, r["claim"][:100], gt,
            r["visual_verdict"],     "YES" if v_ok else "NO", round(r["visual_conf"],     2),
            r["structured_verdict"], "YES" if s_ok else "NO", round(r["structured_conf"],  2),
            r["combined_verdict"],   "YES" if c_ok else "NO", r["resolution_reason"],
        ])
        idx = ws2.max_row
        ws2.cell(idx, 5).fill  = green if v_ok else red
        ws2.cell(idx, 8).fill  = green if s_ok else red
        ws2.cell(idx, 11).fill = green if c_ok else red
    autofit(ws2)

    # ── Sheet 3: Disagreement ─────────────────────────────────────────────────
    ws3 = wb.create_sheet("Disagreement")
    ws3.append(["Case", "Count", "% of total"])
    style_header(ws3)
    total = len(rows)
    labels_d = [
        ("All 3 correct",                    "all_correct"),
        ("Combined wins over both",          "combined_wins_both"),
        ("Combined wins (visual was wrong)", "combined_wins_visual"),
        ("Combined wins (struct was wrong)", "combined_wins_struct"),
        ("Visual only correct",              "visual_only_wins"),
        ("Structured only correct",          "struct_only_wins"),
        ("All 3 wrong",                      "all_wrong"),
    ]
    for label, key in labels_d:
        count = disagree[key]
        ws3.append([label, count, f"{count / total * 100:.1f}%"])
    autofit(ws3)

    wb.save(path)


# ── Main ──────────────────────────────────────────────────────────────────────

async def main(model: str, limit: int, split: str, threshold: float) -> None:
    print(f"\n{'='*65}")
    print(f"  Agent Comparison: Visual vs Structured vs Combined")
    print(f"  {datetime.now().strftime('%Y-%m-%d %H:%M')}")
    print(f"  Model     : {model.split('/')[-1]}")
    print(f"  Samples   : {limit}  |  Split: {split}")
    print(f"  Threshold : {threshold}")
    print("="*65)

    samples = load_samples(split, limit)
    rows: list[dict] = []

    for i, sample in enumerate(samples, 1):
        claim        = sample["claim"].strip()
        ground_truth = sample["label"].strip().upper()
        img_url      = sample["chart_img"].strip()

        print(f"\n[{i}/{len(samples)}] {claim[:70]}...")
        print(f"  GT: {ground_truth}")

        img_path, mime = download_image(img_url)
        if not img_path:
            print("  Skipping: image download failed.")
            continue
        try:
            img_b64 = encode_image(img_path)
        finally:
            remove_temp(img_path)

        t0 = time.time()

        # Agents 2 & 3 run in parallel — same as production
        results = await asyncio.gather(
            run_visual(claim, img_b64, mime, model),
            run_structured(claim, img_b64, mime, model),
            return_exceptions=True,
        )
        vis  = results[0] if not isinstance(results[0], Exception) else None
        strc = results[1] if not isinstance(results[1], Exception) else None
        if isinstance(results[0], Exception): print(f"  Visual error: {results[0]}")
        if isinstance(results[1], Exception): print(f"  Struct error: {results[1]}")

        vis_pred,  vis_conf  = extract_verdict(vis)
        strc_pred, strc_conf = extract_verdict(strc)

        combined_verdict, resolution_reason = await resolve(
            claim, vis, strc, model, threshold)

        elapsed = int((time.time() - t0) * 1000)

        v_mark = "OK" if vis_pred        == ground_truth else "XX"
        s_mark = "OK" if strc_pred       == ground_truth else "XX"
        c_mark = "OK" if combined_verdict == ground_truth else "XX"

        print(f"  Visual     : {vis_pred:<8} {v_mark}  conf={vis_conf:.2f}")
        print(f"  Structured : {strc_pred:<8} {s_mark}  conf={strc_conf:.2f}")
        print(f"  Combined   : {combined_verdict:<8} {c_mark}  [{resolution_reason}]  ({elapsed}ms)")

        rows.append({
            "claim":              claim,
            "ground_truth":       ground_truth,
            "visual_verdict":     vis_pred,
            "visual_conf":        vis_conf,
            "structured_verdict": strc_pred,
            "structured_conf":    strc_conf,
            "combined_verdict":   combined_verdict,
            "resolution_reason":  resolution_reason,
        })

    # ── Metrics ───────────────────────────────────────────────────────────────
    def to_rows(key: str) -> list[dict]:
        return [{"ground_truth": r["ground_truth"], "predicted": r[key]} for r in rows]

    metrics = {
        "visual_only":     compute_metrics(to_rows("visual_verdict")),
        "structured_only": compute_metrics(to_rows("structured_verdict")),
        "combined":        compute_metrics(to_rows("combined_verdict")),
    }
    disagree = disagreement_analysis(rows)

    for label, key in [("Visual Only",     "visual_only"),
                        ("Structured Only", "structured_only"),
                        ("Combined",        "combined")]:
        print_metrics(label, metrics[key])

    n = len(rows)
    print(f"\n{'='*65}")
    print(f"  DISAGREEMENT ANALYSIS  (n={n})")
    print(f"  All 3 correct                    : {disagree['all_correct']}")
    print(f"  Combined wins over BOTH          : {disagree['combined_wins_both']}")
    print(f"  Combined wins (visual wrong)     : {disagree['combined_wins_visual']}")
    print(f"  Combined wins (struct wrong)     : {disagree['combined_wins_struct']}")
    print(f"  Visual only correct              : {disagree['visual_only_wins']}")
    print(f"  Structured only correct          : {disagree['struct_only_wins']}")
    print(f"  All 3 wrong                      : {disagree['all_wrong']}")
    print("="*65)

    # ── Save ──────────────────────────────────────────────────────────────────
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    ts        = datetime.now().strftime("%Y%m%d_%H%M")
    json_path = OUTPUT_DIR / f"agent_comparison_{ts}.json"
    xlsx_path = OUTPUT_DIR / f"agent_comparison_{ts}.xlsx"

    with open(json_path, "w") as f:
        json.dump({
            "model": model, "split": split, "limit": limit,
            "threshold": threshold,
            "metrics": metrics,
            "disagreement": disagree,
            "rows": rows,
        }, f, indent=2)

    save_excel(rows, metrics, disagree, xlsx_path)

    print(f"\n  JSON  -> {json_path.name}")
    print(f"  Excel -> {xlsx_path.name}\n")


if __name__ == "__main__":
    parser = argparse.ArgumentParser(
        description="Compare Visual / Structured / Combined agent strategies",
        formatter_class=argparse.RawDescriptionHelpFormatter,
        epilog=__doc__,
    )
    parser.add_argument("--model",     default=MODEL)
    parser.add_argument("--limit",     type=int,   default=LIMIT)
    parser.add_argument("--split",     default=SPLIT, choices=["test", "test2", "val"])
    parser.add_argument("--threshold", type=float, default=CONFIDENCE_THRESHOLD)
    args = parser.parse_args()

    asyncio.run(main(args.model, args.limit, args.split, args.threshold))
